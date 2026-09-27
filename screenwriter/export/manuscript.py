"""Prose export: compile documents into one manuscript, then PDF, Word or Markdown.

Level-1 headings (# Chapter …) start a new page. [[Notes]] are left out.
A line holding only ***, ---, * * * or # is a scene break.
"""

from __future__ import annotations

import re

from ..marks import strip_marks

NOTE_RE = re.compile(r"\[\[.*?\]\]", re.DOTALL)
SCENE_BREAK_RE = re.compile(r"^\s*(\*\s*\*\s*\*|-{3,}|#)\s*$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")


def clean(text: str) -> str:
    text = strip_marks(NOTE_RE.sub("", text))  # {the hooded figure|Xardas} prints as the phrase
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def compile_markdown(docs: list[tuple[str, str]]) -> str:
    """Join (title, text) documents in order; a document without its own heading
    gets its title as a chapter heading."""
    parts = []
    for title, text in docs:
        body = clean(text)
        if not body:
            continue
        if not body.lstrip().startswith("#"):
            body = f"# {title}\n\n{body}"
        parts.append(body)
    return "\n\n".join(parts) + "\n"


def paragraphs(md: str):
    """Yield (kind, text) with kind in heading1..6, break, quote, bullet, para."""
    for chunk in re.split(r"\n\s*\n", md.strip()):
        lines = chunk.splitlines()
        for line in lines if all(l.lstrip().startswith(("- ", "* ", "#")) for l in lines) else [" ".join(lines)]:
            s = line.strip()
            if SCENE_BREAK_RE.match(s):
                yield "break", ""
            elif m := re.match(r"^(#{1,6})\s+(.*)", s):
                yield f"heading{len(m.group(1))}", m.group(2)
            elif s.startswith(">"):
                yield "quote", re.sub(r"^>\s?", "", s)
            elif s.startswith(("- ", "* ")):
                yield "bullet", s[2:]
            elif s:
                yield "para", s


def runs(text: str):
    """(text, bold, italic) runs for **bold**, *italic* and _italic_."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and len(part) > 4:
            yield part[2:-2], True, False
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            yield part[1:-1], False, True
        else:
            yield part, False, False


def write_markdown(md: str, path: str) -> None:
    with open(path, "w", encoding="utf-8") as f:
        f.write(md)


def write_docx(md: str, path: str, title: str = "") -> None:
    """Standard manuscript format: Times 12pt, double spaced, 1" margins,
    first-line indents, chapters on new pages."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Inches, Pt

    doc = Document()
    doc.core_properties.title = title
    for section in doc.sections:
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    first_block = True
    after_heading = False
    for kind, text in paragraphs(md):
        if kind.startswith("heading"):
            p = doc.add_paragraph()
            if kind == "heading1" and not first_block:
                p.paragraph_format.page_break_before = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(144 if kind == "heading1" else 24)
            p.paragraph_format.space_after = Pt(24)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(14 if kind == "heading1" else 12)
            after_heading = True
        elif kind == "break":
            p = doc.add_paragraph("#")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            after_heading = True
        else:
            p = doc.add_paragraph()
            if kind == "quote":
                p.paragraph_format.left_indent = Inches(0.5)
            elif kind == "bullet":
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.add_run("•\t")
            elif not after_heading:
                p.paragraph_format.first_line_indent = Inches(0.5)
            for s, bold, italic in runs(text):
                run = p.add_run(s)
                run.bold = bold
                run.italic = italic or kind == "quote"
            after_heading = False
        first_block = False
    doc.save(path)


def write_pdf(md: str, path: str, paper: str = "Letter", title: str = "") -> int:
    """Typeset with Qt's rich text engine; returns the page count."""
    from PySide6.QtCore import QMarginsF, QRectF, QSizeF, Qt
    from PySide6.QtGui import (
        QFont,
        QPageLayout,
        QPageSize,
        QPainter,
        QPdfWriter,
        QTextBlockFormat,
        QTextCursor,
        QTextDocument,
        QTextFormat,
    )

    from ..editors.common import first_available_font

    size = QPageSize(QPageSize.PageSizeId.A4 if paper == "A4" else QPageSize.PageSizeId.Letter)
    writer = QPdfWriter(path)
    writer.setPageLayout(QPageLayout(size, QPageLayout.Orientation.Portrait, QMarginsF(72, 72, 72, 72)))
    writer.setResolution(72)
    writer.setTitle(title)
    writer.setCreator("Screenwriter")

    doc = QTextDocument()
    font = QFont(first_available_font("Iowan Old Style", "Charter", "Georgia", "Times New Roman"))
    font.setPointSizeF(12)
    doc.setDefaultFont(font)
    doc.setMarkdown(md)

    first = True
    after_heading = False
    block = doc.begin()
    while block.isValid():
        cursor = QTextCursor(block)
        fmt = QTextBlockFormat(block.blockFormat())
        level = fmt.headingLevel()
        text = block.text().strip()
        if level:
            if level == 1 and not first:
                fmt.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
            fmt.setAlignment(Qt.AlignmentFlag.AlignHCenter)
            fmt.setTopMargin(96 if level == 1 else 18)
            fmt.setBottomMargin(24)
            after_heading = True
        elif SCENE_BREAK_RE.match(text) or fmt.hasProperty(QTextFormat.Property.BlockTrailingHorizontalRulerWidth):
            fmt.setAlignment(Qt.AlignmentFlag.AlignHCenter)
            after_heading = True
        else:
            fmt.setLineHeight(150, QTextBlockFormat.LineHeightTypes.ProportionalHeight.value)
            fmt.setTopMargin(0)
            fmt.setBottomMargin(0)
            if not after_heading and not block.textList() and fmt.indent() == 0 and fmt.leftMargin() == 0:
                fmt.setTextIndent(24)
            after_heading = False
        cursor.setBlockFormat(fmt)
        first = False
        block = block.next()

    # Paginate ourselves so we control page numbers (bottom centre, from page 2).
    width, height = size.sizePoints().width() - 144, size.sizePoints().height() - 144
    doc.setPageSize(QSizeF(width, height))
    painter = QPainter(writer)
    number_font = QFont(font)
    number_font.setPointSizeF(10)
    pages = doc.pageCount()
    for page in range(pages):
        if page:
            writer.newPage()
        painter.save()
        painter.translate(0, -page * height)
        doc.drawContents(painter, QRectF(0, page * height, width, height))
        painter.restore()
        if page:
            painter.setFont(number_font)
            painter.drawText(QRectF(0, height + 24, width, 20), Qt.AlignmentFlag.AlignHCenter, str(page + 1))
    painter.end()
    return pages
