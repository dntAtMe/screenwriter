import re
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest

from screenwriter.export import manuscript, screenplay
from screenwriter.export.screenplay import MORE, El, blocks, clean, paginate

SAMPLE = Path(__file__).parent.parent / "examples" / "The Lighthouse"


def test_clean_line():
    assert clean(".flashback [[check]]", El.SCENE) == "FLASHBACK"
    assert clean("She **really** runs _fast_.", El.ACTION) == "She really runs fast."
    assert clean("> THE END <", El.CENTERED) == "THE END"
    assert clean("@McCLANE ^", El.CHARACTER) == "McCLANE"


def test_blocks_group_speech_and_skip_notes():
    fields, script = blocks("Title: X\n\n# ACT\n\n= synopsis\n\nINT. A - DAY\n\nMARA\n(quiet)\nHi.\n\n[[note]]\n\nShe goes.\nAnd goes.")
    assert fields == {"title": "X"}
    assert [(b.kind, [e for e, _ in b.lines]) for b in script] == [
        (El.SCENE, [El.SCENE]),
        (El.CHARACTER, [El.CHARACTER, El.PARENTHETICAL, El.DIALOGUE]),
        (El.ACTION, [El.ACTION, El.ACTION]),
    ]


def long_script(scenes=30):
    speech = " ".join(["This line of dialogue goes on and on so it wraps."] * 12)
    out = []
    for i in range(scenes):
        out.append(f"INT. ROOM {i} - DAY\n\nSomeone enters. " + "They look around. " * 10)
        out.append(f"MARA\n(quietly)\n{speech}")
        out.append("OWEN\nShort answer.")
    return "\n\n".join(out)


def test_pagination_rules():
    _, pages = paginate(long_script(), 54)
    assert len(pages) > 5
    for page in pages:
        assert len(page) <= 54
        printed = [l for l in page if l is not None]
        assert page[0] is not None and page[-1] is not None  # no leading/trailing blanks
        assert printed[-1][0] != El.SCENE  # heading kept with what follows
        assert printed[-1][0] != El.PARENTHETICAL
    flat = [l for page in pages for l in page if l is not None]
    assert (El.CHARACTER, MORE) in flat
    assert (El.CHARACTER, "MARA (CONT'D)") in flat
    # every page after a (MORE) opens with the CONT'D cue
    for a, b in zip(pages, pages[1:]):
        if a[-1] == (El.CHARACTER, MORE):
            assert b[0][1].endswith("(CONT'D)")


def test_wrapping_widths():
    _, pages = paginate(long_script(3), 54)
    for page in pages:
        for line in page:
            if line:
                el, s = line
                assert len(s) <= screenplay.COLUMNS[el][1]


def test_fdx():
    root = ET.fromstring(screenplay.to_fdx((SAMPLE / "docs" / "pilot.fountain").read_text()))
    paras = root.findall("./Content/Paragraph")
    types = [p.get("Type") for p in paras]
    assert types[0] == "Scene Heading" and "Character" in types and "Dialogue" in types and "Parenthetical" in types
    assert paras[0].findtext("Text") == "EXT. SKERRY ROCK LIGHTHOUSE - DUSK"
    assert root.find("./TitlePage/Content/Paragraph/Text").text == "THE LIGHTHOUSE"


def test_screenplay_pdf(qapp, tmp_path):
    out = tmp_path / "pilot.pdf"
    pages = screenplay.write_pdf(long_script(), str(out))
    assert pages > 5
    assert out.read_bytes().startswith(b"%PDF")


def test_compile_and_paragraphs():
    md = manuscript.compile_markdown([("One", "# Chapter One\n\nText [[note]] here."), ("Two", "No heading.\n\n***\n\n> quote")])
    assert md == "# Chapter One\n\nText  here.\n\n# Two\n\nNo heading.\n\n***\n\n> quote\n"
    assert list(manuscript.paragraphs(md)) == [
        ("heading1", "Chapter One"), ("para", "Text  here."), ("heading1", "Two"),
        ("para", "No heading."), ("break", ""), ("quote", "quote"),
    ]
    assert list(manuscript.runs("a **b** *c* _d_")) == [
        ("a ", False, False), ("b", True, False), (" ", False, False), ("c", False, True),
        (" ", False, False), ("d", False, True),
    ]


def test_docx(tmp_path):
    from docx import Document

    md = manuscript.compile_markdown([("One", "# One\n\nFirst *para*.\n\nSecond."), ("Two", "# Two\n\nThird.")])
    out = tmp_path / "book.docx"
    manuscript.write_docx(md, str(out), "Book")
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert texts == ["One", "First para.", "Second.", "Two", "Third."]
    assert doc.paragraphs[3].paragraph_format.page_break_before
    assert doc.paragraphs[1].runs[1].italic


def test_manuscript_pdf(qapp, tmp_path):
    md = manuscript.compile_markdown([("One", "# One\n\n" + "Words " * 800), ("Two", "# Two\n\nMore.")])
    out = tmp_path / "book.pdf"
    assert manuscript.write_pdf(md, str(out)) >= 3
    assert out.read_bytes().startswith(b"%PDF")


def test_export_formats_for_sample(qapp, tmp_path):
    from screenwriter.exportdialog import export_formats
    from screenwriter.project import Project

    project = Project.open(SAMPLE)
    text_of = project.read_text
    what, formats = export_formats(project, project.find("manuscript"), text_of)
    assert "2 prose documents" in what and [f.extension for f in formats] == ["pdf", "docx", "md"]
    what, formats = export_formats(project, project.find("pilot"), text_of)
    assert [f.extension for f in formats] == ["pdf", "fdx", "fountain"]
    what, formats = export_formats(project, project.find("storymap"), text_of)
    png = tmp_path / "map.png"
    formats[0].run(str(png), "Letter")
    assert png.read_bytes().startswith(b"\x89PNG")
    what, formats = export_formats(project, project.find("screenplay"), text_of)  # folder with no prose
    assert formats == []


def pdf_pages(path):
    """(text, inked pixels) for each page, rendered the way a viewer would: on white paper."""
    from PySide6.QtCore import QSize, Qt
    from PySide6.QtGui import QImage, QPainter
    from PySide6.QtPdf import QPdfDocument

    doc = QPdfDocument()
    doc.load(str(path))
    pages = []
    for i in range(doc.pageCount()):
        image = doc.render(i, QSize(306, 396))
        paper = QImage(image.size(), QImage.Format.Format_RGB32)
        paper.fill(Qt.GlobalColor.white)
        painter = QPainter(paper)
        painter.drawImage(0, 0, image)
        painter.end()
        ink = sum(1 for y in range(0, paper.height(), 2) for x in range(0, paper.width(), 2)
                  if paper.pixelColor(x, y).lightness() < 160)
        pages.append((doc.getAllText(i).text(), ink))
    doc.close()
    return pages


@pytest.mark.parametrize("node_id, expected", [
    ("pilot", ["THE LIGHTHOUSE", "EXT. SKERRY ROCK LIGHTHOUSE - DUSK", "MARA", "It's Thursday."]),
    ("ch01", ["Chapter 1 — The Keeper", "Skerry Rock"]),
    ("manuscript", ["Chapter 1 — The Keeper", "Chapter 2 — Fog"]),
    ("storymap", ["Lighthouse", "Mara Quinn"]),  # Qt drops the "Th" ligature from a board's text layer
])
def test_pdf_pages_have_content(qapp, tmp_path, node_id, expected):
    """Every exported page shows something, and the text is really in the PDF (not blank pages)."""
    from screenwriter.exportdialog import export_formats
    from screenwriter.project import Project

    project = Project.open(SAMPLE)
    _, formats = export_formats(project, project.find(node_id), project.read_text)
    out = tmp_path / f"{node_id}.pdf"
    next(f for f in formats if f.extension == "pdf").run(str(out), "Letter")
    pages = pdf_pages(out)
    assert pages
    assert all(ink > 0 for _, ink in pages), [ink for _, ink in pages]
    from PySide6.QtGui import QFontDatabase

    if not QFontDatabase.families():
        pytest.skip("no fonts for headless Qt here: the PDF has boxes instead of text")
    squeeze = lambda s: "".join(s.split())  # a board's text comes out a letter per line
    text = squeeze(" ".join(t for t, _ in pages))
    for s in expected:
        assert squeeze(s) in text
    assert "[[" not in text  # notes are never printed


def test_wrap_line_matches_textwrap():
    import random
    import textwrap

    from screenwriter.export.screenplay import wrap_line

    rng = random.Random(7)
    words = ["a", "keeper", "lighthouse", "x" * 45, "—", "on", "the", "rocks,", "(beat)", "supercalifragilistic"]
    for _ in range(500):
        s = " ".join(rng.choice(words) for _ in range(rng.randint(0, 30)))
        s = " " * rng.randint(0, 2) + s + " " * rng.randint(0, 2)
        for width in (25, 35, 38, 60):
            assert wrap_line(s, width) == textwrap.wrap(s, width, break_on_hyphens=False), (s, width)


def test_emphasis_is_kept_as_styles():
    from screenwriter.export.screenplay import BOLD, ITALIC, UNDERLINE, runs

    s = clean(r"She **really** runs _fast_, *very* ***very*** _**far**_ \*ok\*.", El.ACTION)
    assert s == "She really runs fast, very very far *ok*."
    assert [(t, st) for _, t, st in runs(s) if st] == [
        ("really", BOLD), ("fast", UNDERLINE), ("very", ITALIC), ("very", BOLD | ITALIC), ("far", BOLD | UNDERLINE),
    ]
    # styles survive the forcing mark, capitals and wrapping
    heading = clean(".**flashback**", El.SCENE)
    assert heading == "FLASHBACK" and set(heading.styles) == {BOLD}
    _, pages = paginate("INT. A - DAY\n\n" + "plain " * 12 + "**" + ("loud " * 12).strip() + "**", 54)
    lines = [line[1] for line in pages[0] if line and line[0] == El.ACTION]
    assert len(lines) > 1 and all(len(line.styles) == len(line) for line in lines)
    assert lines[-1].styles[-1] == BOLD and lines[0].styles[0] == 0


def test_fdx_keeps_styles():
    root = ET.fromstring(screenplay.to_fdx("INT. A - DAY\n\nShe **really** _runs_."))
    texts = [(t.text, t.get("Style")) for t in root.findall("./Content/Paragraph")[1].findall("Text")]
    assert texts == [("She ", None), ("really", "Bold"), (" ", None), ("runs", "Underline"), (".", None)]


def test_pdf_prints_bold_and_italic(qapp, tmp_path):
    out = tmp_path / "styled.pdf"
    screenplay.write_pdf("INT. A - DAY\n\nPlain **bold** *italic* _under_.", str(out))
    data = out.read_bytes()
    fonts = set(re.findall(rb"/BaseFont\s*/([^\s/>]+)", data))
    assert len(fonts) >= 3, fonts  # regular, bold and italic faces
    assert any(b"Bold" in f for f in fonts) and any(b"Italic" in f or b"Oblique" in f for f in fonts), fonts
